"""Pure parsing helpers: decoding, delimiter/header detection, tolerant number/date parsing, fuzzy mapping."""

from __future__ import annotations

import csv
import difflib
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

PREVIEW_ROWS = 20
MAX_ROWS = 20000


@dataclass
class ParsedTable:
    headers: list[str]
    rows: list[list[str]]
    encoding: str = "utf-8"
    delimiter: str = ","
    header_row: int = 0
    sheet: str | None = None
    warnings: list[str] = field(default_factory=list)


def decode(raw: bytes) -> tuple[str, str]:
    """Return (text, encoding). BOM wins; then strict UTF-8; else Windows-1252."""
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw[3:].decode("utf-8", errors="replace"), "utf-8-sig"
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16"), "utf-16"
    try:
        return raw.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace"), "windows-1252"


def detect_delimiter(text: str) -> str:
    sample = "\n".join(text.splitlines()[:30])
    try:
        return csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in (";", ",", "\t", "|")}
        best = max(counts, key=counts.get)  # type: ignore[arg-type]
        return best if counts[best] else ","


def _is_numberish(cell: str) -> bool:
    return bool(re.fullmatch(r"[-+()\d.,\s%€$]+", cell)) and any(ch.isdigit() for ch in cell)


def detect_header_row(rows: list[list[str]]) -> int:
    """First of the top 10 rows that is about as wide as the data and mostly text."""
    top = rows[:10]
    if not top:
        return 0
    width = max(sum(1 for c in r if c.strip()) for r in top)
    for i, r in enumerate(top):
        filled = [c for c in r if c.strip()]
        if len(filled) >= max(1, -(-width * 3 // 5)) and sum(not _is_numberish(c) for c in filled) >= len(filled) * 0.8:
            return i
    return 0


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _finish(grid: list[list[str]], table: ParsedTable) -> ParsedTable:
    grid = [r for r in grid if any(c.strip() for c in r)]
    hi = detect_header_row(grid)
    headers = [h.strip() for h in grid[hi]] if grid else []
    seen: dict[str, int] = {}
    for i, h in enumerate(headers):
        h = h or f"Column {i + 1}"
        if h in seen:
            seen[h] += 1
            h = f"{h} ({seen[h]})"
        else:
            seen[h] = 1
        headers[i] = h
    body = grid[hi + 1 :]
    if len(body) > MAX_ROWS:
        table.warnings.append(f"Only the first {MAX_ROWS} rows are imported")
        body = body[:MAX_ROWS]
    n = len(headers)
    table.headers = headers
    table.rows = [(r + [""] * n)[:n] for r in body]
    table.header_row = hi
    return table


def parse_csv(raw: bytes) -> ParsedTable:
    text, enc = decode(raw)
    delim = detect_delimiter(text)
    grid = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), delimiter=delim)]
    return _finish(grid, ParsedTable([], [], encoding=enc, delimiter=delim))


def parse_xlsx(raw: bytes) -> ParsedTable:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    grid = [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    return _finish(grid, ParsedTable([], [], encoding="xlsx", delimiter="", sheet=ws.title))


def parse_file(filename: str, raw: bytes) -> ParsedTable:
    if filename.lower().endswith((".xlsx", ".xlsm")) or raw[:4] == b"PK\x03\x04":
        return parse_xlsx(raw)
    return parse_csv(raw)


# ── Values ────────────────────────────────────────────────────────────────
def parse_decimal(value: str) -> Decimal:
    """Accept 1.234,56 / 1,234.56 / 1 234,56 / 1234,5 / (12,00) / 12,00- / €12."""
    s = (value or "").strip()
    if not s:
        raise ValueError("empty")
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg, s = True, s[1:-1]
    if s.endswith("-"):
        neg, s = True, s[:-1]
    s = re.sub(r"[\s'  €$£%]|EUR|USD|CHF|GBP", "", s, flags=re.I)
    if s.startswith("-"):
        neg, s = not neg, s[1:]
    if "," in s and "." in s:
        dec = "," if s.rfind(",") > s.rfind(".") else "."
        s = s.replace("." if dec == "," else ",", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", "") if re.fullmatch(r"\d{1,3}(,\d{3}){2,}", s) else s.replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    if not re.fullmatch(r"\d+(\.\d+)?|\.\d+", s):
        raise ValueError(f"not a number: {value!r}")
    try:
        d = Decimal(s)
    except InvalidOperation as exc:
        raise ValueError(f"not a number: {value!r}") from exc
    return -d if neg else d


def parse_date(value: str) -> date:
    """dd.mm.yyyy, dd.mm.yy, yyyy-mm-dd, dd/mm/yyyy, dd-mm-yyyy (day first unless impossible)."""
    s = (value or "").strip().split(" ")[0].split("T")[0]
    if not s:
        raise ValueError("empty")
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        return date(int(m[1]), int(m[2]), int(m[3]))
    m = re.fullmatch(r"(\d{1,2})[./-](\d{1,2})[./-](\d{2}|\d{4})", s)
    if not m:
        raise ValueError(f"not a date: {value!r}")
    a, b, y = int(m[1]), int(m[2]), int(m[3])
    if y < 100:
        y += 2000 if y < 70 else 1900
    if b > 12 and a <= 12:
        a, b = b, a
    return date(y, b, a)


EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[A-Za-z]{2,}$")


def normalize_vat(value: str) -> str:
    return re.sub(r"[\s.\-]", "", value or "").upper()


def is_valid_vat(value: str) -> bool:
    v = normalize_vat(value)
    if v.startswith("DE"):
        return bool(re.fullmatch(r"DE\d{9}", v))
    if v.startswith("ATU"):
        return bool(re.fullmatch(r"ATU\d{8}", v))
    if v.startswith("CHE"):
        return bool(re.fullmatch(r"CHE\d{9}(MWST|TVA|IVA)?", v))
    return bool(re.fullmatch(r"[A-Z]{2}[0-9A-Z+*]{2,13}", v))


# ── Header matching ───────────────────────────────────────────────────────
def norm_header(h: str) -> str:
    s = unicodedata.normalize("NFKD", h or "").replace("ß", "ss")
    s = "".join(ch for ch in s if not unicodedata.combining(ch)).lower()
    return re.sub(r"[^0-9a-zа-яё]", "", s)


def auto_map(headers: list[str], synonyms: dict[str, list[str]], threshold: float = 0.82) -> dict[str, str]:
    """Map header -> field. Exact synonym hits first, then fuzzy; each field used once."""
    return {h: f for h, (f, _s) in auto_map_scored(headers, synonyms, threshold).items()}


def auto_map_scored(
    headers: list[str], synonyms: dict[str, list[str]], threshold: float = 0.82
) -> dict[str, tuple[str, float]]:
    """Like auto_map but returns header -> (field, confidence 0..1)."""
    norm_syn = {f: {norm_header(x) for x in xs + [f]} for f, xs in synonyms.items()}
    cands: list[tuple[float, int, str, str]] = []
    for idx, h in enumerate(headers):
        nh = norm_header(h)
        if not nh:
            continue
        for f, syns in norm_syn.items():
            if nh in syns:
                score = 1.0
            else:
                score = max(difflib.SequenceMatcher(None, nh, s).ratio() for s in syns)
            if score >= threshold:
                cands.append((score, -idx, h, f))
    out: dict[str, tuple[str, float]] = {}
    used: set[str] = set()
    for score, _i, h, f in sorted(cands, reverse=True):
        if h not in out and f not in used:
            out[h] = (f, round(score, 2))
            used.add(f)
    return out
